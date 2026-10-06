#!/usr/bin/env python
"""Fill episode_meta CSVs in Excel with drop-down lists taken from configs/episode_meta_schema.yaml.

The CSV in episode_meta/ stays the record; the .xlsx is a scratch copy for typing.

    uv run python scripts/episode_meta_xlsx.py export episode_meta/<dataset>.csv    # -> outputs/episode_meta_xlsx/<dataset>.xlsx
    (fill it in Excel, save as .xlsx)
    uv run python scripts/episode_meta_xlsx.py import outputs/episode_meta_xlsx/<dataset>.xlsx   # -> back into the CSV

Drop-downs follow the schema: `strict` fields reject anything outside `values` (Excel "stop");
other fields with `values` only warn. `multi` fields (mechanism) take several values separated
by ';' -- Excel cannot validate that, so the cell only warns and `import` checks every value.
`import` refuses to write while a strict or required check fails.
"""

import argparse
import csv
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parents[1]
SCHEMA = REPO / "configs" / "episode_meta_schema.yaml"
XLSX_DIR = REPO / "outputs" / "episode_meta_xlsx"


def load_fields(schema: Path) -> dict[str, dict]:
    return {f["name"]: f for f in yaml.safe_load(schema.read_text(encoding="utf-8"))["fields"]}


def export(csv_path: Path, schema: Path, xlsx_dir: Path = XLSX_DIR) -> Path:
    fields = load_fields(schema)
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    header, body = rows[0], rows[1:]
    wb = Workbook()
    ws = wb.active
    ws.title = "episodes"
    ws.append(header)
    for r in body:  # episode_index as a number ("0.0" from pandas-written CSVs -> 0); everything else as text
        ws.append([int(float(v)) if c == "episode_index" and v else v for c, v in zip(header, r)])
    # value lists live on a hidden sheet, so lists longer than Excel's 255-char inline limit still work
    lists = wb.create_sheet("lists")
    lists.sheet_state = "hidden"
    n = len(body) + 1
    for col_idx, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        f = fields.get(name)
        if not f:
            continue
        if f.get("required"):
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        tip = f"{f.get('prompt', name)}\n{' '.join(str(f.get('help', '')).split())}"
        cell.comment = Comment(tip[:1000], "schema")
        values = [str(v) for v in f.get("values") or [] if str(v) != ""]
        if not values:
            continue
        letter = cell.column_letter
        for i, v in enumerate(values, start=1):
            lists[f"{letter}{i}"] = v
        strict = f.get("strict") and not f.get("multi")
        dv = DataValidation(
            type="list",
            formula1=f"=lists!${letter}$1:${letter}${len(values)}",
            allow_blank=not f.get("required"),
            showErrorMessage=True,
            errorStyle="stop" if strict else "warning",
            errorTitle=name,
            error=("Only: " if strict else "Not in the list (ok if intended): ") + ", ".join(values)
            + ("  -- several values: separate with ;" if f.get("multi") else ""),
            promptTitle=name[:32],
            prompt=(", ".join(values) + ("  (; for several)" if f.get("multi") else ""))[:255],
            showInputMessage=True,
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{n}")
    ws.freeze_panes = "B2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(10, min(28, max(len(str(c.value or "")) for c in col) + 2))
    xlsx_dir.mkdir(parents=True, exist_ok=True)
    out = xlsx_dir / (csv_path.stem + ".xlsx")
    if out.exists():
        raise SystemExit(f"{out} exists -- import it first, or delete it to start over from the CSV")
    wb.save(out)
    return out


def import_(xlsx_path: Path, schema: Path, csv_dir: Path) -> Path:
    fields = load_fields(schema)
    ws = load_workbook(xlsx_path, data_only=True)["episodes"]
    rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    header, body = rows[0], [r for r in rows[1:] if any(r)]
    errors, warnings = [], []
    for r in body:
        ep = r[header.index("episode_index")] if "episode_index" in header else "?"
        for name, v in zip(header, r):
            f = fields.get(name)
            if not f:
                continue
            allowed = {str(x) for x in f.get("values") or []} | {""}
            if f.get("required") and v == "":
                errors.append(f"episode {ep}: {name} is required")
            parts = [p.strip() for p in v.split(";")] if f.get("multi") else [v]
            bad = [p for p in parts if p and f.get("values") and p not in allowed]
            if bad and f.get("strict"):
                errors.append(f"episode {ep}: {name} {bad} not in {sorted(allowed - {''})}")
            elif bad:  # non-strict: one summary line per field, not one per episode
                warnings.append((name, tuple(sorted(allowed - {""}))))
            if f.get("multi") and v:
                r[header.index(name)] = ";".join(p for p in parts if p)
    for (name, suggested), count in sorted({w: warnings.count(w) for w in warnings}.items()):
        print(f"~ {name}: {count} row(s) outside the suggested values {list(suggested)} (allowed, not strict)")
    for e in errors:
        print("!", e)
    if errors:
        raise SystemExit(f"{len(errors)} problem(s) -- fix them in {xlsx_path} and import again; the CSV was not touched")
    out = csv_dir / (xlsx_path.stem + ".csv")
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([header] + body)
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("mode", choices=["export", "import"])
    ap.add_argument("path", type=Path)
    ap.add_argument("--schema", type=Path, default=SCHEMA)
    ap.add_argument("--csv-dir", type=Path, default=REPO / "episode_meta", help="import: where the CSV is written")
    ap.add_argument("--xlsx-dir", type=Path, default=XLSX_DIR, help="export: where the .xlsx is written")
    a = ap.parse_args()
    if a.mode == "export":
        print("wrote", export(a.path, a.schema, a.xlsx_dir))
    else:
        print("wrote", import_(a.path, a.schema, a.csv_dir))


if __name__ == "__main__":
    main()
